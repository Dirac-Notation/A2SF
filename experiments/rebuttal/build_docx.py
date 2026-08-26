"""REBUTTAL_DRAFT_KO.md -> REBUTTAL_DRAFT_KO.docx (주간 보고서 스타일).
마크다운 초안을 파싱해서 Word 문서 생성: 제목/헤딩/불릿/표.
  python experiments/rebuttal/build_docx.py
"""
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
import sys
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "REBUTTAL_DRAFT_KO.md")
OUT = SRC.replace(".md", ".docx")


def set_korean_font(doc):
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")


def strip_md(text):
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text


def add_table(doc, header, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for j, h in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = strip_md(h)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for i, row in enumerate(rows):
        for j in range(len(header)):
            t.rows[i + 1].cells[j].text = strip_md(row[j]) if j < len(row) else ""
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = "맑은 고딕"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")


def main():
    doc = Document()
    set_korean_font(doc)

    lines = open(SRC).read().split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        # 표 감지
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[-| ]+\|?$", lines[i + 1].strip()):
            header = [c.strip() for c in line.strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, header, rows)
            doc.add_paragraph()
            continue
        if line.startswith("# "):
            doc.add_heading(strip_md(line[2:]), level=0)
        elif line.startswith("## "):
            doc.add_heading(strip_md(line[3:]), level=1)
        elif line.startswith("### "):
            doc.add_heading(strip_md(line[4:]), level=2)
        elif line.strip() == "---":
            pass
        elif line.startswith("  - "):
            doc.add_paragraph(strip_md(line.strip()[2:]), style="List Bullet 2")
        elif line.startswith("- "):
            doc.add_paragraph(strip_md(line[2:]), style="List Bullet")
        elif line.strip():
            doc.add_paragraph(strip_md(line))
        i += 1

    doc.save(OUT)
    print(f"saved -> {OUT}")


if __name__ == "__main__":
    main()
